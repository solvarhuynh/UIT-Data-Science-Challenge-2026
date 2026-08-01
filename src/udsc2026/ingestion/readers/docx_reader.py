"""Reader for Microsoft Word documents."""

from pathlib import Path
from typing import Any, Dict, Iterable, List

from udsc2026.ingestion.readers._helpers import first_nonempty_line, make_doc_id
from udsc2026.ingestion.readers.models import RawDocument


def _require_docx() -> Any:
    try:
        from docx import Document
    except ImportError as error:  # pragma: no cover - incomplete install only
        raise ImportError(
            "Reading DOCX files requires python-docx. "
            "Install it with `pip install python-docx`."
        ) from error
    return Document


def _table_lines(table: Any) -> Iterable[str]:
    """Yield table text; cells are retained as rows rather than discarded."""
    for row in table.rows:
        cells: List[str] = []
        for cell in row.cells:
            cells.append(_cell_text(cell))
        yield "\t".join(cells)


def _cell_text(cell: Any) -> str:
    """Collect paragraph and nested-table content from a table cell in order."""
    parts: List[str] = []
    for block in cell.iter_inner_content():
        if hasattr(block, "rows"):
            parts.extend(_table_lines(block))
        else:
            parts.append(block.text)
    return "\n".join(parts)


def _document_lines(document: Any) -> Iterable[str]:
    """Yield top-level paragraphs and tables in their source-document order."""
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document).text
        elif isinstance(child, CT_Tbl):
            yield from _table_lines(Table(child, document))


def read_docx(file_path: str) -> RawDocument:
    """Read paragraphs and tables from a DOCX source without text cleaning."""
    document = _require_docx()(file_path)
    path = Path(file_path)

    lines: List[str] = list(_document_lines(document))
    raw_text = "\n".join(lines)

    core_properties = document.core_properties
    title = (core_properties.title or "").strip() or first_nonempty_line(raw_text)
    metadata: Dict[str, Any] = {
        "file_name": path.name,
        "file_size_bytes": path.stat().st_size,
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
    }
    if core_properties.author:
        metadata["author"] = core_properties.author
    if core_properties.created:
        metadata["created"] = core_properties.created.isoformat()
    if core_properties.modified:
        metadata["modified"] = core_properties.modified.isoformat()

    return RawDocument(
        doc_id=make_doc_id(str(path)),
        source_path=str(path),
        title=title,
        raw_text=raw_text,
        file_format="docx",
        metadata=metadata,
    )
